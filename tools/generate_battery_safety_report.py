from pathlib import Path
from zipfile import ZipFile
import shutil

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT_DIR = Path("dist")
OUT_DIR.mkdir(parents=True, exist_ok=True)
OUT = OUT_DIR / "1.docx"
ALIAS = OUT_DIR / "面向搭载锂离子电池特种装备的电池大模型及主动安全防护技术研究指南及论证报告.docx"

NAVY = "17365D"
TEAL = "0F6B78"
WHITE = "FFFFFF"
DARK = "222222"
GRAY = "666666"
FONT_HEAD = "Microsoft YaHei"
FONT_BODY = "SimSun"
TITLE = "面向搭载锂离子电池特种装备的电池大模型及主动安全防护技术研究"

GUIDE_GOAL = (
    "针对具身智能、引信装备、储能装备等搭载锂离子电池的特种装备存在健康状态评估维度单一、主动安全模型跨材料体系、跨车型和跨场景迁移泛化能力弱，"
    "以及特种装备自身数据规模小、故障标签稀缺等问题，开展电池多维健康指数、跨域时间序列预训练模型、故障预警与根因分析智能体研究，突破多源异构电池数据治理、"
    "容量—电压—电阻—故障四类健康指数构建、跨材料体系与跨场景自监督预训练、少样本迁移、故障风险识别、根因分析和运维建议生成等关键技术，研制电池大模型及主动安全防护软件原理样机，"
    "实现跨三元与磷酸铁锂体系、跨30种以上车型及车载—储能场景的状态评估和风险预警，为搭载锂离子电池的特种装备提供可迁移、可解释、可追溯的健康管理和主动安全防护技术支撑。"
)

GUIDE_CONTENTS = [
    "1）电池多维健康指数与运维看板。构建容量型、电压型、电阻型、故障型四类健康指数，形成不少于50项指标的定义、算法、入口条件、证据和版本，研究多维特征赋权与综合评分方法。",
    "2）电池时间序列预训练模型。利用车载和储能场景大规模时序数据开展掩码重构、下一时刻预测、对比学习和多任务预训练，形成跨材料体系、跨车型、跨充电/换电模式和跨场景的统一表征与迁移方法。",
    "3）故障预警大模型与主动安全防护智能体。融合多维健康指数、预训练潜在特征、故障规则、维修标签和相似案例，形成故障风险识别、Top-K根因分析、维修保养建议和人工复核闭环。",
]

GUIDE_INDICATORS = [
    "1）形成容量型、电压型、电阻型、故障型四类电池健康指数及运维看板，健康指数数量不少于50项；形成基于多维特征赋权的综合健康评分。",
    "2）构建电池时间序列预训练模型，测试集覆盖三元、磷酸铁锂等多种材料体系，覆盖不少于30种中国主流新能源车型并包含充电、换电等模式及典型真实储能场景；95%的场景单元电压还原误差不高于30 mV，全部场景单元SOH平均绝对误差不高于5个百分点。",
    "3）构建故障预警大模型和电池状态评估及故障预警智能体；对经维修或复检确认的故障电池，故障查准率不低于85%，输出Top-K根因、证据摘要及维修保养建议。",
]

GUIDE_OUTPUTS = (
    "电池大模型及主动安全防护软件原理样机1套；多维电池健康指数体系与运维看板1套；时间序列预训练模型及迁移适配工具包1套；"
    "故障预警与根因分析智能体1套；经评审冻结的测试大纲、测试集说明和测试报告1套；研究报告及技术规范若干。"
)

BACKGROUND = (
    "大量具身智能、引信装备、储能装备等特种装备使用锂离子电池，但其电池数据规模小、运行工况差异大，普遍存在健康状态评估维度不足、故障机理难解释、主动安全模型迁移泛化能力弱等问题。"
    "实验室和车载场景已积累较丰富的电压、电流、温度、SOC、故障码、维修和容量检测数据，但传统小模型通常依赖特定电芯、车型和工况，难以直接迁移至数据稀缺的特种装备。"
    "现有电池健康评价多依赖SOH或单一阈值，不能同时反映容量衰减、电压一致性、内阻增长和故障隐患；故障样本稀缺又限制了监督模型训练。"
    "项目拟利用车载和储能大规模无标签数据开展时间序列预训练，以少量目标域样本完成适配，并将多维健康指数、预训练潜在表征、故障规则和维修结果融合为主动安全防护能力，为特种装备电池提供统一、可迁移和可审计的状态评估与风险预警底座。"
)

RESEARCH_GOAL = (
    "针对搭载锂离子电池特种装备健康状态难以全面评价、故障风险难以及时识别、模型跨材料体系和跨场景泛化能力不足等需求，开展多维电池健康指数、时间序列预训练模型和故障预警智能体研究，"
    "突破多源时序数据治理、健康指标物理语义映射、跨车型与跨化学体系自监督表征、少样本域适配、故障根因分析和证据治理等关键技术，研制电池大模型及主动安全防护软件原理样机，"
    "形成四类不少于50项健康指数、覆盖30种以上车型和典型储能场景的预训练模型、故障预警及根因分析智能体，实现高质量电池状态评估、故障预警、维修保养建议和全过程可追溯，"
    "为具身智能、引信装备、储能装备等搭载电池设备的主动安全防护提供技术支撑。"
)

RESEARCH_ITEMS = [
    ("多源数据治理与冻结测试集", "建立车辆、电池包、电芯、储能站点、时间窗口和维修结果的统一标识与数据契约，完成去噪、对齐、工况分段、单位统一、标签分级和数据血缘；按对象与时间隔离训练、验证和测试数据，防止数据泄漏。"),
    ("四类电池健康指数与综合评分", "围绕容量、电压、电阻和故障四类维度，构建不少于50项可计算健康指数，明确对象层级、输入窗口、公式、入口条件、缺失处理、正常范围、证据和版本；研究物理规则与数据驱动相结合的赋权和综合评分。"),
    ("跨域时间序列预训练模型", "采用Transformer等时序编码器，通过掩码重构、未来片段预测、跨周期对比和多任务学习吸收大规模车载与储能无标签数据；研究材料体系编码、车型路由、域适配、参数高效微调、适用域识别和安全回退。"),
    ("故障预警与根因分析智能体", "融合健康指数、潜在时序特征、故障规则、维修标签和相似案例，输出故障概率、风险分级、Top-K根因及证据；通过智能体编排数据检查、模型路由、结果复核、维修保养建议和运维工单草案。"),
]

KEY_TECHS = [
    ("多维健康指数的物理语义与赋权", "将容量衰减、电压曲线与一致性、内阻与极化、故障事件和潜在特征转化为可解释指标，利用客观权重、专家权重和场景权重形成综合健康评分，并保留每项评分的贡献来源。"),
    ("跨材料体系与跨场景时序预训练", "以海量无标签序列学习电池正常演化和异常偏离的通用表示，通过材料体系、车型和场景标识、域适配及少样本微调，使模型从车载数据迁移到储能和特种装备场景。"),
    ("模型路由、适用域与可信回退", "根据材料体系、车型、采样频率、数据窗口和质量状态选择预训练模型、小模型、规则或物理基线；低置信或适用域外样本进入人工复核，避免长尾场景静默失败。"),
    ("多模型融合的故障预警与根因分析", "联合健康指数、Transformer潜在特征、异常检测、故障树、相似案例和维修反馈，形成风险评分与Top-K根因；使用维修或复检确认结果作为最终验收标签。"),
]

INDICATORS = [
    ("1", "多维健康指数与看板", "形成容量型、电压型、电阻型、故障型四类健康指数及运维看板，健康指数数量≥50项；全部指标具备定义、计算逻辑、入口条件、数据质量状态、证据来源和版本。", "T-01：健康指数与运维看板测试"),
    ("2", "时间序列预训练与泛化", "测试集覆盖三元、磷酸铁锂等材料体系，≥30种中国主流新能源车型，包含充电、换电及典型真实储能场景；95%的场景单元电压还原误差≤30 mV，全部场景单元SOH平均绝对误差≤5个百分点。", "T-02：预训练模型与跨域泛化测试"),
    ("3", "故障预警与智能体", "在覆盖≥30种车型及典型真实储能场景的冻结测试集上，对维修或复检确认的故障电池，故障查准率≥85%；智能体输出Top-K根因、证据摘要和维修保养建议。", "T-03：故障预警与智能体测试"),
]

EXPECTED_EFFECT = (
    "形成覆盖容量、电压、电阻和故障四维度的电池健康评价方法，建立跨材料体系、跨30种以上车型和车载—储能场景的时间序列预训练模型，提升数据稀缺装备场景的模型迁移能力；"
    "形成故障预警、根因分析和维修建议智能体，使电池状态评价由单一指标转向全维度、可解释、可追溯的主动安全防护，为具身智能、引信装备和储能装备提供统一技术底座。"
)

FOUNDATION_PARAGRAPHS = [
    "在电池系统层面，申请人将目光拓展至智能电池管理与优化。他牵头开发了基于大数据和大模型的电池管理系统，以车辆及储能系统运行数据为基础（千亿条真实用户电池场景数据），融合机器学习和智能决策算法，实现了电池全生命周期的健康状态评估与优化控制[Energy Storage Mater. 2026, 86, 104897；申请欧洲发明专利EP4693592A1、申请中国发明专利202410175617.0]。荣获挑战杯2025年度中国青年科技创新“揭榜挂帅”擂台赛——基于AI大模型的新能源汽车动力电池安全预警与健康监测技术赛道特等奖（前十获奖团队中唯一高校团队）。上述工作为本项目构建跨场景电池状态评估、模型评测和安全预警闭环提供了丰富的工程化落地经验。",
    "申请人善于利用大语言模型等先进AI技术对接复杂系统的大规模数据和机理模型。例如，他构建了多智能体协作的储能材料创意生成平台，用于新机理探索和新材料开发；在内部多任务评测集中，所开发多智能体系统在若干材料研发子任务上表现出较强的问题分解、工具调用和证据整合能力，公开基准验证仍在持续开展。团队还开发了用于自主识别意图、自主生成分析方案、自主研究并发掘新机制的低温电导率智能机器人，获得第十一届全国大学生物理实验竞赛（创新）一等奖、第十八届北京市大学生物理实验竞赛一等奖。这些基础可直接支撑本项目的模型路由、智能体编排、Critical评测和证据治理。",
]

REFERENCES = [
    "[F1] 《电池基因库及电池事故溯源分析平台构建方案》：全生命周期数据、数字档案、潜在特征和健康指数体系。",
    "[F2] 《新能源汽车电池热失控事故溯源分析大模型解读》：Transformer/LSTM/VAE、自监督预训练、多任务微调、故障根因及边云协同。",
    "[F3] 《申报人情况概述-北航-王诗童 V2.0》：申请人与团队研究基础。",
    "[W1] NIST, Artificial Intelligence Risk Management Framework (AI RMF 1.0), 2023。",
]


def set_font(run, name=FONT_BODY, size=10.5, bold=None, color=DARK):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    for key in ("w:eastAsia", "w:ascii", "w:hAnsi"):
        rpr.rFonts.set(qn(key), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell(cell, text, size=9, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, color=DARK, font_name=FONT_BODY):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.06
    r = p.add_run(str(text))
    set_font(r, font_name, size, bold, color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_borders(table, color="666666", size="6"):
    tblpr = table._tbl.tblPr
    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def add_table(doc, headers, rows, widths, font_size=8.5, header_size=8.8):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_borders(table)
    for i, h in enumerate(headers):
        set_cell(table.rows[0].cells[i], h, header_size, True, WD_ALIGN_PARAGRAPH.CENTER, WHITE, FONT_HEAD)
        shade_cell(table.rows[0].cells[i], NAVY)
        table.rows[0].cells[i].width = Cm(widths[i])
    for ri, row in enumerate(rows):
        rr = table.add_row()
        for i, value in enumerate(row):
            set_cell(rr.cells[i], value, font_size, False, WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT)
            if ri % 2 == 1:
                shade_cell(rr.cells[i], "F8F9FA")
            rr.cells[i].width = Cm(widths[i])
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.add_run(text)
    return p


def add_body(doc, text, indent=True, size=10.5):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.32
    p.paragraph_format.space_after = Pt(4)
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    r = p.add_run(text)
    set_font(r, FONT_BODY, size, False, DARK)
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])
    set_font(run, FONT_HEAD, 8.5, False, GRAY)


doc = Document()
doc.core_properties.title = TITLE + "指南及论证报告"
doc.core_properties.author = "北京航空航天大学 王诗童"

normal = doc.styles["Normal"]
normal.font.name = FONT_BODY
normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
normal.font.size = Pt(10.5)
for level, size, color in [(1, 15, NAVY), (2, 12, TEAL)]:
    style = doc.styles[f"Heading {level}"]
    style.font.name = FONT_HEAD
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_HEAD)
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)

# Attachment 1: guide table
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width = Cm(29.7)
sec.page_height = Cm(21)
sec.top_margin = Cm(0.8)
sec.bottom_margin = Cm(0.8)
sec.left_margin = Cm(0.6)
sec.right_margin = Cm(0.6)

p = doc.add_paragraph(); set_font(p.add_run("附件1："), FONT_BODY, 10.5)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; set_font(p.add_run("指南建议"), FONT_HEAD, 15, True)
main_text = (
    "研究目标：" + GUIDE_GOAL + "\n研究内容：\n" + "\n".join(GUIDE_CONTENTS)
    + "\n主要技术指标：\n" + "\n".join(GUIDE_INDICATORS) + "\n预期成果：" + GUIDE_OUTPUTS
)
headers = ["序号", "指南名称", "研究目标、内容和主要技术指标", "研究周期", "经费需求\n（万元）", "编制单位", "建议人", "密级"]
rows = [["1", TITLE, main_text, "2年", "200", "北京航空航天大学", "姓名：王诗童\n手机号：18811361360", "按申报要求填写"]]
add_table(doc, headers, rows, [1.0, 3.7, 13.9, 1.6, 1.7, 2.1, 2.5, 1.5], 6.7, 7.2)

# Attachment 2: report
sec2 = doc.add_section(WD_SECTION.NEW_PAGE)
sec2.orientation = WD_ORIENT.PORTRAIT
sec2.page_width = Cm(21)
sec2.page_height = Cm(29.7)
sec2.top_margin = Cm(1.7)
sec2.bottom_margin = Cm(1.6)
sec2.left_margin = Cm(2.0)
sec2.right_margin = Cm(2.0)
sec2.footer.is_linked_to_previous = False
add_page_number(sec2.footer.paragraphs[0])

p = doc.add_paragraph(); set_font(p.add_run("附件2："), FONT_BODY, 10.5)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; set_font(p.add_run(TITLE + "论证报告"), FONT_HEAD, 17, True)

add_heading(doc, "一、研究背景")
add_body(doc, BACKGROUND)
add_heading(doc, "二、研究目标")
add_body(doc, RESEARCH_GOAL)
add_heading(doc, "三、研究内容")
for i, (title, desc) in enumerate(RESEARCH_ITEMS, 1):
    add_body(doc, f"{i}）{title}：{desc}", False)
add_heading(doc, "四、拟突破的关键技术")
for i, (title, desc) in enumerate(KEY_TECHS, 1):
    add_body(doc, f"{i}）{title}：{desc}", False)
add_heading(doc, "五、技术指标")
add_table(doc, ["序号", "指标", "考核要求", "对应测试"], INDICATORS, [1.0, 3.2, 8.0, 4.0], 8.2, 8.6)
add_heading(doc, "六、预期成效")
add_body(doc, EXPECTED_EFFECT)
add_heading(doc, "七、成果形式及验证考虑")
outputs = [
    ("电池大模型及主动安全防护软件原理样机", "包含健康指数与看板、预训练模型与迁移、故障预警与智能体、模型评测与证据治理等功能。", "软件功能演示、安装包、用户手册和测试报告。"),
    ("多维电池健康指数体系", "容量型、电压型、电阻型、故障型四类不少于50项健康指数及综合评分方法。", "指标字典、算法说明、示例数据和运维看板测试。"),
    ("时间序列预训练模型及迁移工具包", "覆盖三元、磷酸铁锂、30种以上车型及车载—储能场景的预训练模型、适配器、模型卡和评测脚本。", "冻结测试集、模型复测和跨域评测报告。"),
    ("故障预警与根因分析智能体", "风险识别、Top-K根因、证据摘要、维修保养建议和人工复核工作流。", "冻结故障测试集、维修真值核验和端到端任务回放。"),
    ("测试与研究报告", "经评审冻结的测试大纲、测试报告、技术研究报告、数据规范和部署运维规范。", "专家评审、第三方或项目组测试。"),
]
add_table(doc, ["成果形式", "主要内容", "验证方式"], outputs, [4.0, 7.5, 4.7], 8.3, 8.7)

add_heading(doc, "7.1 测试大纲建议", 2)
tests = [
    ("T-01 健康指数与看板", "冻结指标字典、公式、样例数据和软件版本。", "逐项检查四类指标、缺失处理、评分贡献和看板显示。", "健康指数≥50项，分类正确，公式可复算，综合评分可追溯。"),
    ("T-02 预训练模型与泛化", "冻结数据授权、车型/站点清单、材料体系、训练/验证/测试划分、真值和评价脚本。", "按场景单元离线运行，计算电压还原误差和SOH误差，并检查对象与时间隔离。", "95%场景单元电压误差≤30 mV；全部场景单元SOH MAE≤5个百分点；无测试泄漏。"),
    ("T-03 故障预警与智能体", "冻结维修或复检确认的正负样本、≥30车型和典型储能场景、报告模板。", "运行风险识别、根因分析和建议生成，计算查准率并抽查证据。", "故障查准率≥85%；输出Top-K根因、证据和维修保养建议；低置信进入人工复核。"),
]
add_table(doc, ["测试项", "测试前置条件", "测试步骤", "通过判据"], tests, [3.5, 4.3, 5.5, 3.0], 7.9, 8.3)

add_heading(doc, "八、研制周期和经费需求")
add_body(doc, "研究周期为24个月，建议经费200万元。经费主要用于多源数据治理、预训练算力、模型开发与评测、健康指数和看板开发、智能体与软件集成、第三方测试、专家咨询及部署运维，不购置大型固定设备。")
phases = [
    ("第一阶段", "第1—4个月", "冻结数据、任务、指标和测试大纲；完成数据清单、场景单元定义、健康指标候选和人工基线。"),
    ("第二阶段", "第5—10个月", "完成数据治理、四类健康指数、综合评分、车载与储能预训练数据集和预训练模型V1。"),
    ("第三阶段", "第11—18个月", "完成跨车型/化学体系/场景迁移、故障预警大模型、根因分析智能体和运维看板。"),
    ("第四阶段", "第19—24个月", "冻结测试集并正式测试，完成问题整改、第三方或专家评审、软件与报告交付。"),
]
add_table(doc, ["阶段", "周期", "主要任务"], phases, [3.0, 3.5, 9.7], 8.5, 8.8)

add_heading(doc, "九、已有基础")
for paragraph in FOUNDATION_PARAGRAPHS:
    add_body(doc, paragraph)
add_heading(doc, "十、编制单位和人员")
add_body(doc, "编制单位：北京航空航天大学", False)
add_body(doc, "姓名：王诗童", False)
add_body(doc, "手机号码：18811361360", False)
add_heading(doc, "主要资料依据", 2)
for reference in REFERENCES:
    add_body(doc, reference, False, 9.0)

doc.save(OUT)
shutil.copy2(OUT, ALIAS)
with ZipFile(OUT) as package:
    assert package.testzip() is None
    assert "word/document.xml" in package.namelist()
print(f"Generated {OUT} ({OUT.stat().st_size} bytes)")
